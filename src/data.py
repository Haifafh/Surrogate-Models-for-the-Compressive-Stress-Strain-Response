"""
Dataset construction for the nanosilica/PVA-ECC compressive stress--strain benchmark.

Source data
-----------
The raw measurements are the digitised compressive stress--strain curves of 24
engineered cementitious composite (ECC) mixtures, published alongside:

    Albadrani, M.A. et al. Polymers 2025, 17, 2067.

Workbook layout ("curves_source_workbook.xlsx"):
    sheets 'a'..'f'  -> nanosilica content nS = 0, 1, 2, 3, 4, 5 % (by mass of binder)
    columns N1..N4   -> PVA fibre volume fraction 0.5, 1.0, 1.5, 2.0 %
    each pair (strain, stress): strain in units of 1e3 microstrain, stress in MPa

Provenance caveat
-----------------
The digitised peak strengths differ systematically from the values tabulated in
Table 5 of the companion study (mean absolute deviation 13.0 %, up to 27.3 %).
The digitised curves agree with the characteristic strengths quoted in the text
and abstract of that study. See docs/REPRODUCIBILITY.md and the manuscript's
Section 2.2. All models in this repository are trained and compared on the
single consistent digitised dataset shipped here.
"""
import os
import numpy as np
import pandas as pd

SHEETS = ['a', 'b', 'c', 'd', 'e', 'f']
PVA_LEVELS = [0.5, 1.0, 1.5, 2.0]
GRID_STEP_UM = 100.0            # common axial-strain grid, microstrain
MODULUS_FRACTION = 0.4          # E0 fitted on the branch below 0.4*sigma_peak

# Five complete mixtures withheld in the blind protocol (Split B), as (nS %, PVA %)
TEST_COMBOS = [(1.0, 0.5), (2.0, 1.5), (3.0, 2.0), (4.0, 1.0), (5.0, 1.5)]

REPO_ROOT = os.environ.get('ECC_REPO', os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RAW_XLSX = os.environ.get('ECC_RAW_XLSX', os.path.join(REPO_ROOT, 'data', 'raw', 'curves_source_workbook.xlsx'))


def load_raw_curves(path=None):
    """Parse the source workbook into a long-form frame: nS, PVA, strain_um_m, stress_MPa."""
    import openpyxl
    path = path or RAW_XLSX
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []
    for si, sheet in enumerate(SHEETS):
        ws = wb[sheet]
        for ci, pva in enumerate(PVA_LEVELS):
            col_e, col_s = 2 * ci + 1, 2 * ci + 2      # 1-based columns
            for r in range(2, ws.max_row + 1):
                e = ws.cell(r, col_e).value
                s = ws.cell(r, col_s).value
                try:
                    e = float(e); s = float(s)
                except (TypeError, ValueError):
                    continue
                rows.append((float(si), float(pva), e * 1000.0, s))
    return pd.DataFrame(rows, columns=['nS', 'PVA', 'strain_um_m', 'stress_MPa'])


def resample(df_raw, step=GRID_STEP_UM):
    """Linearly interpolate every curve onto a common 100-microstrain grid."""
    out = []
    for (ns, pva), g in df_raw.groupby(['nS', 'PVA'], sort=True):
        g = g.sort_values('strain_um_m')
        e = g['strain_um_m'].to_numpy(float)
        s = g['stress_MPa'].to_numpy(float)
        keep = np.concatenate(([True], np.diff(e) > 0))     # strictly increasing
        e, s = e[keep], s[keep]
        grid = np.arange(0.0, np.floor(e[-1] / step) * step + 0.5 * step, step)
        grid = grid[grid <= e[-1]]
        out.append(pd.DataFrame({
            'nS': ns, 'PVA': pva,
            'strain_um_m': grid,
            'strain_pct': grid / 1e4,
            'eps_frac': grid / 1e6,
            'stress_MPa': np.interp(grid, e, s),
        }))
    return pd.concat(out, ignore_index=True)


def peak_properties(raw, grid):
    """Peak stress/strain and initial modulus E0, evaluated on the RAW digitised points.

    Peaks and E0 are computed before interpolation (a resampled grid can clip an
    off-grid peak); the point count and maximum strain are reported from the
    resampled grid, matching the manuscript's Table 1.
    """
    rows = []
    for (ns, pva), g in raw.groupby(['nS', 'PVA'], sort=True):
        g = g.sort_values('strain_um_m')
        e = g['strain_um_m'].to_numpy(float)
        s = g['stress_MPa'].to_numpy(float)
        keep = np.concatenate(([True], np.diff(e) > 0))
        e, s = e[keep], s[keep]
        k = int(np.argmax(s))
        fc, e0 = float(s[k]), float(e[k])
        asc = (s < MODULUS_FRACTION * fc) & (e <= e0)
        E0 = float(np.polyfit(e[asc], s[asc], 1)[0] * 1000.0) if asc.sum() >= 2 else np.nan
        gg = grid[(grid.nS == ns) & (grid.PVA == pva)]
        rows.append(dict(nS=ns, PVA=pva, sigma_peak=fc, strain_peak_um_m=e0,
                         strain_peak_pct=e0 / 1e4, E0_GPa=E0,
                         n_points=int(len(gg)), strain_max_um_m=float(gg.strain_um_m.max())))
    return pd.DataFrame(rows)


def build_dataset(raw_path=None):
    """Full digitisation pipeline: raw workbook -> tidy curves on the 100-microstrain grid."""
    return resample(load_raw_curves(raw_path))


def split_A(df):
    """Random point-wise 80/20 partition (interpolation within mixtures already seen)."""
    rng = np.random.default_rng(0)
    idx = rng.permutation(len(df))
    n_tr = int(round(0.8 * len(df)))
    return df.iloc[idx[:n_tr]].copy(), df.iloc[idx[n_tr:]].copy()


def split_B(df, test_combos=None):
    """Blind-mixture partition: withhold five complete stress--strain curves."""
    test_combos = test_combos or TEST_COMBOS
    key = list(zip(df['nS'].round(6), df['PVA'].round(6)))
    mask = np.array([k in set((round(a, 6), round(b, 6)) for a, b in test_combos) for k in key])
    return df.loc[~mask].copy(), df.loc[mask].copy()


def load_dataset(csv_path=None):
    """Load the shipped, resampled tidy dataset."""
    csv_path = csv_path or os.path.join(REPO_ROOT, 'data', 'dataset_curves.csv')
    return pd.read_csv(csv_path)
